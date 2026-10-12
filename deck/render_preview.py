"""Approximate PPTX -> PNG renderer for visual QA (no office suite in sandbox).

Draws fills, outlines, pictures, simple autoshapes and wrapped text with DejaVu.
Not pixel-accurate vs PowerPoint — good enough to catch layout mistakes.
"""
import math
import os
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

PX = 120.0  # px per inch
REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "preview")
os.makedirs(OUT, exist_ok=True)

_fontcache = {}


def font(size_pt, bold):
    key = (round(size_pt, 1), bold)
    if key not in _fontcache:
        _fontcache[key] = ImageFont.truetype(BOLD if bold else REG,
                                             max(6, int(round(size_pt * PX / 72.0))))
    return _fontcache[key]


def emu2px(v):
    return v / 914400.0 * PX


def rgb_tuple(c):
    return (c[0], c[1], c[2])


def fill_of(sp):
    try:
        f = sp.fill
        if f.type is not None and int(f.type) == 1:  # solid
            col = rgb_tuple(bytes(f.fore_color.rgb))
            alpha = 1.0
            sfill = f._xPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
            if sfill is not None:
                clr = sfill.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
                if clr is not None:
                    a = clr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')
                    if a is not None:
                        alpha = int(a.get('val')) / 100000.0
            return col, alpha
    except Exception:
        pass
    return None, 1.0


def line_of(sp):
    try:
        ln = sp.line
        if ln.fill.type is not None and int(ln.fill.type) == 1:
            w = ln.width
            return rgb_tuple(bytes(ln.color.rgb)), (w / 12700.0 if w else 1.0)
    except Exception:
        pass
    return None, 1.0


def rot_pts(pts, cx, cy, deg):
    if not deg:
        return pts
    a = math.radians(deg)
    ca, sa = math.cos(a), math.sin(a)
    out = []
    for x, y in pts:
        dx, dy = x - cx, y - cy
        out.append((cx + dx * ca - dy * sa, cy + dx * sa + dy * ca))
    return out


def draw_poly(img, pts, fill, alpha, line, lw):
    layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
    dr = ImageDraw.Draw(layer)
    if fill:
        dr.polygon(pts, fill=fill + (int(255 * alpha),))
    if line:
        dr.line(pts + [pts[0]], fill=line + (255,), width=max(1, int(round(lw * PX / 72.0))))
    img.paste(Image.alpha_composite(Image.new("RGBA", img.size, (0, 0, 0, 0)), layer), (0, 0),
              layer)


def draw_ellipse(img, box, fill, alpha, line, lw, deg=0, cx=None, cy=None):
    layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
    dr = ImageDraw.Draw(layer)
    dr.ellipse(box, fill=(fill + (int(255 * alpha),)) if fill else None,
               outline=(line + (255,)) if line else None,
               width=max(1, int(round(lw * PX / 72.0))) if line else 0)
    if deg:
        layer = layer.rotate(-deg, center=(cx, cy), resample=Image.BICUBIC)
    img.paste(layer, (0, 0), layer)


def arrow_pts(kind, x, y, w, h):
    if kind == MSO_SHAPE.RIGHT_ARROW:
        hw = h * 0.55 / 2 * 2 * 0.5
        head_h = h
        shaft_h = h * 0.34
        hy = y + (h - shaft_h) / 2
        return [(x, hy), (x + w - head_h * 0.6, hy), (x + w - head_h * 0.6, y),
                (x + w, y + h / 2), (x + w - head_h * 0.6, y + h),
                (x + w - head_h * 0.6, y + h - (h - shaft_h) / 2 - (h - shaft_h) / 2 + (h - shaft_h) / 2),
                (x, y + h - (h - shaft_h) / 2)]
    if kind == MSO_SHAPE.DOWN_ARROW:
        shaft_w = w * 0.34
        hx = x + (w - shaft_w) / 2
        return [(hx, y), (hx + shaft_w, y), (hx + shaft_w, y + h - w * 0.6),
                (x + w, y + h - w * 0.6), (x + w / 2, y + h), (x, y + h - w * 0.6),
                (hx, y + h - w * 0.6)]
    if kind == MSO_SHAPE.LEFT_RIGHT_ARROW:
        sh = h * 0.34
        hy = y + (h - sh) / 2
        aw = w * 0.22
        return [(x, y + h / 2), (x + aw, y), (x + aw, hy), (x + w - aw, hy), (x + w - aw, y),
                (x + w, y + h / 2), (x + w - aw, y + h), (x + w - aw, y + h - sh / 1 - (h - sh) + (h - sh)),
                (x + aw, y + hy - y + h - sh - (hy - y)), (x + aw, y + h)]
    return [(x, y), (x + w, y), (x + w, y + h), (x, y + h)]


def wrap(text, fnt, maxw, dr):
    words = text.split()
    lines, cur = [], ""
    for wd in words:
        t = (cur + " " + wd).strip()
        if dr.textlength(t, font=fnt) <= maxw or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = wd
    if cur:
        lines.append(cur)
    return lines or [""]


def draw_text(sp, img, dr, x, y, w, h):
    tf = sp.text_frame
    anchor = tf.vertical_anchor
    paras = []
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            continue
        words = []
        maxsize = 12
        for r in runs:
            size = r.font.size.pt if r.font.size else 12
            bold = bool(r.font.bold)
            try:
                col = rgb_tuple(bytes(r.font.color.rgb)) if (r.font.color and
                                                             r.font.color.type is not None) \
                    else (0, 0, 0)
            except Exception:
                col = (0, 0, 0)
            f = font(size, bold)
            maxsize = max(maxsize, size)
            parts = r.text.split(" ")
            for i, wd in enumerate(parts):
                if i > 0:
                    words.append((" ", f, col, size))
                if wd:
                    words.append((wd, f, col, size))
        paras.append((words, maxsize, p.alignment, p.line_spacing or 1.0,
                      p.space_after.pt if p.space_after else 0))
    if not paras:
        return
    inset = (0.0 if sp.shape_type == MSO_SHAPE_TYPE.TEXT_BOX else 0.08) * PX
    maxw = w - 2 * inset
    blocks = []
    total = 0.0
    for words, maxsize, al, ls, sa in paras:
        lh = maxsize * ls * 1.22 * PX / 72.0
        lines, cur, curw = [], [], 0.0
        for wd, f, col, size in words:
            ww = dr.textlength(wd, font=f)
            if wd == " ":
                if cur:
                    cur.append((wd, f, col))
                    curw += ww
                continue
            if curw + ww > maxw and cur:
                lines.append(cur)
                cur, curw = [], 0.0
                while wd == " ":
                    break
            cur.append((wd, f, col))
            curw += ww + dr.textlength(" ", font=f)
        if cur:
            lines.append(cur)
        blocks.append((lines, al, lh, sa * PX / 72.0))
        total += lh * len(lines) + sa * PX / 72.0
    cy = y + inset
    if anchor == MSO_ANCHOR.MIDDLE:
        cy = y + h / 2 - total / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        cy = y + h - total
    for lines, al, lh, sa in blocks:
        for ln in lines:
            tw = sum(dr.textlength(wd, font=f) for wd, f, _ in ln)
            if al == PP_ALIGN.CENTER:
                tx = x + w / 2 - tw / 2
            elif al == PP_ALIGN.RIGHT:
                tx = x + w - inset - tw
            else:
                tx = x + inset
            for wd, f, col in ln:
                dr.text((tx, cy), wd, font=f, fill=col)
                tx += dr.textlength(wd, font=f) + dr.textlength(" ", font=f)
            cy += lh
        cy += sa


def render(pptx_path):
    prs = Presentation(pptx_path)
    W = int(emu2px(prs.slide_width))
    H = int(emu2px(prs.slide_height))
    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGBA", (W, H), (255, 255, 255, 255))
        dr = ImageDraw.Draw(img)
        for sp in slide.shapes:
            if sp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                x, y = emu2px(sp.left), emu2px(sp.top)
                w, h = emu2px(sp.width), emu2px(sp.height)
                try:
                    pim = Image.open(__import__("io").BytesIO(sp.image.blob)).convert("RGBA")
                    pim = pim.resize((max(1, int(w)), max(1, int(h))))
                    img.paste(pim, (int(x), int(y)), pim)
                except Exception as e:
                    print("pic fail", e)
                continue
            if sp.shape_type == MSO_SHAPE_TYPE.LINE or sp.__class__.__name__ == "Connector":
                try:
                    x1, y1 = emu2px(sp.begin_x), emu2px(sp.begin_y)
                    x2, y2 = emu2px(sp.end_x), emu2px(sp.end_y)
                    col, lw = line_of(sp)
                    dr.line([(x1, y1), (x2, y2)], fill=(col or (0, 0, 0)) + (255,),
                            width=max(1, int(lw * PX / 72.0)))
                except Exception:
                    pass
                continue
            x, y = emu2px(sp.left), emu2px(sp.top)
            w, h = emu2px(sp.width), emu2px(sp.height)
            fill, alpha = fill_of(sp)
            line, lw = line_of(sp)
            deg = sp.rotation or 0
            cx, cy = x + w / 2, y + h / 2
            try:
                ast = sp.auto_shape_type
            except Exception:
                ast = None
            if ast == MSO_SHAPE.OVAL:
                draw_ellipse(img, [x, y, x + w, y + h], fill, alpha, line, lw, deg, cx, cy)
            elif ast in (MSO_SHAPE.RIGHT_ARROW, MSO_SHAPE.DOWN_ARROW, MSO_SHAPE.LEFT_RIGHT_ARROW):
                pts = rot_pts(arrow_pts(ast, x, y, w, h), cx, cy, deg)
                draw_poly(img, pts, fill, alpha, line, lw)
            elif ast == MSO_SHAPE.ISOSCELES_TRIANGLE:
                pts = rot_pts([(x + w / 2, y), (x + w, y + h), (x, y + h)], cx, cy, deg)
                draw_poly(img, pts, fill, alpha, line, lw)
            elif ast == MSO_SHAPE.CROSS:
                t = w / 3
                pts = [(x + t, y), (x + 2 * t, y), (x + 2 * t, y + t), (x + w, y + t),
                       (x + w, y + 2 * t), (x + 2 * t, y + 2 * t), (x + 2 * t, y + h),
                       (x + t, y + h), (x + t, y + 2 * t), (x, y + 2 * t), (x, y + t),
                       (x + t, y + t)]
                draw_poly(img, rot_pts(pts, cx, cy, deg), fill, alpha, line, lw)
            elif ast == MSO_SHAPE.CHEVRON:
                t = w * 0.3
                pts = [(x, y), (x + w - t, y), (x + w, y + h / 2), (x + w - t, y + h),
                       (x, y + h), (x + t, y + h / 2)]
                draw_poly(img, rot_pts(pts, cx, cy, deg), fill, alpha, line, lw)
            elif ast in (MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.ROUND_2_SAME_RECTANGLE):
                r = 0.12 * min(w, h)
                try:
                    r = sp.adjustments[0] * min(w, h)
                except Exception:
                    pass
                layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
                d2 = ImageDraw.Draw(layer)
                d2.rounded_rectangle([x, y, x + w, y + h], radius=r,
                                     fill=(fill + (int(255 * alpha),)) if fill else None,
                                     outline=(line + (255,)) if line else None,
                                     width=max(1, int(round(lw * PX / 72.0))) if line else 0)
                if deg:
                    layer = layer.rotate(-deg, center=(cx, cy), resample=Image.BICUBIC)
                img.paste(layer, (0, 0), layer)
            elif fill or line:
                pts = rot_pts([(x, y), (x + w, y), (x + w, y + h), (x, y + h)], cx, cy, deg)
                draw_poly(img, pts, fill, alpha, line, lw)
            if sp.has_text_frame and sp.text_frame.text.strip():
                draw_text(sp, img, dr, x, y, w, h)
        img.convert("RGB").save(os.path.join(OUT, f"slide{idx:02d}.png"))
        print("rendered", idx)


if __name__ == "__main__":
    render(sys.argv[1] if len(sys.argv) > 1 else
           os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "CareRoute_Hackathon_Deck_TeamTechHead.pptx"))
