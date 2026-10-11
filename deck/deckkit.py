"""deckkit — tiny design system for building the CareRoute hackathon deck (python-pptx).

Canvas: 13.333in x 7.5in (16:9). All coordinates in inches.
Palette (exact, from the brief):
  ELECTRIC  #0719C8   ORANGE #FFAA16   MAGENTA #F3265B   PURPLE #683C9E
  MINT      #58BE98   YELLOW #FFD21A   BLACK   #080808   WHITE  #FFFFFF
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
import math

# ---------------------------------------------------------------- palette
ELECTRIC = "0719C8"
ORANGE   = "FFAA16"
MAGENTA  = "F3265B"
PURPLE   = "683C9E"
MINT     = "58BE98"
YELLOW   = "FFD21A"
BLACK    = "080808"
WHITE    = "FFFFFF"

# soft tints used only for card fills on white slides (derived from palette)
BLUE_TINT   = "EEF1FB"
ORANGE_TINT = "FFF6E5"
MAGENTA_TINT= "FDEBF0"
PURPLE_TINT = "F2EDF8"
MINT_TINT   = "EAF7F1"
YELLOW_TINT = "FFFBE6"
GREY_LINE   = "D8DCE8"
INK_SOFT    = "3A3F52"   # secondary body text on light surfaces (near-black, readable)

HEAD = "Segoe UI"
BODY = "Segoe UI"

SW, SH = 13.333, 7.5
M = 0.55                      # page margin

# approx Segoe UI advance widths (em fraction)
_CW = {"reg": 0.505, "bold": 0.545}


def rgb(h):
    return RGBColor.from_string(h)


def _para_lines(text, size, bold, width_in):
    """greedy word-wrap; returns number of visual lines."""
    if not text:
        return 1
    cw = _CW["bold" if bold else "reg"] * size / 72.0
    avail = max(width_in, 0.2)
    lines, cur = 1, 0.0
    for word in text.split():
        w = (len(word) + (1 if cur > 0 else 0)) * cw
        if cur + w > avail and cur > 0:
            lines += 1
            cur = len(word) * cw
        else:
            cur += w
    return lines


# ---------------------------------------------------------------- deck
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.slide = None
        self._n = 0
        self.warnings = []

    # -- slides -----------------------------------------------------------
    def new_slide(self):
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self._n += 1
        return self.slide

    def notes(self, text):
        self.slide.notes_slide.notes_text_frame.text = text

    # -- primitives -------------------------------------------------------
    def shape(self, kind, x, y, w, h, fill=None, line=None, lw=1.0,
              dash=None, rad=None, rot=0.0, shadow=False):
        sp = self.slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = rgb(fill)
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = rgb(line)
            sp.line.width = Pt(lw)
            if dash:
                sp.line.dash_style = dash
        if rad is not None and kind in (MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.ROUND_1_RECTANGLE):
            try:
                sp.adjustments[0] = rad
            except Exception:
                pass
        if rot:
            sp.rotation = rot
        sp.shadow.inherit = shadow
        if sp.has_text_frame:
            sp.text_frame.word_wrap = True
        return sp

    def rect(self, x, y, w, h, fill=None, **kw):
        return self.shape(MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, **kw)

    def rrect(self, x, y, w, h, fill=None, rad=0.10, **kw):
        return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, rad=rad, **kw)

    def oval(self, x, y, w, h, fill=None, **kw):
        return self.shape(MSO_SHAPE.OVAL, x, y, w, h, fill=fill, **kw)

    def pill(self, x, y, w, h, fill=None, **kw):
        return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, rad=0.5, **kw)

    def picture(self, path, x, y, w, h):
        return self.slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))

    # -- text -------------------------------------------------------------
    def text(self, x, y, w, h, paras, anchor="top", align="left", wrap=True,
             check=True, name=None):
        """paras: list of dicts or list-of-runs.
        para dict keys: t (str) or runs (list of run dicts), size, color, bold,
        italic, align, space_before, space_after, line (spacing multiple), font."""
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                              "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        total_h = 0.0
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            runs = p.get("runs") if isinstance(p, dict) else p
            if runs is None:
                runs = [{"t": p.get("t", "")}]
            base = p if isinstance(p, dict) else {}
            palign = p.get("align", align) if isinstance(p, dict) else align
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                              "right": PP_ALIGN.RIGHT}[palign]
            ls = base.get("line", 1.0)
            para.line_spacing = ls
            para.space_before = Pt(base.get("space_before", 0))
            para.space_after = Pt(base.get("space_after", 0))
            sizes = []
            sample = ""
            anybold = False
            for r in runs:
                run = para.add_run()
                run.text = r.get("t", "")
                f = run.font
                f.size = Pt(r.get("size", base.get("size", 12)))
                f.bold = r.get("bold", base.get("bold", False))
                f.italic = r.get("italic", base.get("italic", False))
                f.name = r.get("font", base.get("font", BODY))
                f.color.rgb = rgb(r.get("color", base.get("color", BLACK)))
                sizes.append(r.get("size", base.get("size", 12)))
                anybold = anybold or r.get("bold", base.get("bold", False))
                sample += r.get("t", "")
            maxsize = max(sizes) if sizes else 12
            # height estimate
            lines = _para_lines(sample, maxsize, anybold, w) if wrap else 1
            total_h += (lines * maxsize * ls * 1.22 + base.get("space_before", 0)
                        + base.get("space_after", 0)) / 72.0
        if check and total_h > h + 0.06:
            self.warnings.append(
                f"slide {self._n}: text overflow ~{total_h:.2f}in in {h:.2f}in box :: "
                f"{(name or paras[0].get('t', runs[0].get('t',''))[:48] if isinstance(paras[0], dict) else '')!r}")
        return tb

    def chip(self, x, y, w, h, t, fill, color, size=9.5, bold=True, line=None, lw=1.0):
        self.pill(x, y, w, h, fill=fill, line=line, lw=lw)
        self.text(x, y, w, h, [{"t": t, "size": size, "color": color, "bold": bold,
                                "align": "center"}], anchor="middle", check=False)

    # -- arrows -----------------------------------------------------------
    def harrow(self, x, y, w, color, th=0.16):
        sp = self.shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, th, fill=color)
        try:
            sp.adjustments[0] = 0.55
            sp.adjustments[1] = 0.34
        except Exception:
            pass
        return sp

    def varrow(self, x, y, h, color, th=0.16):
        sp = self.shape(MSO_SHAPE.DOWN_ARROW, x, y, th, h, fill=color)
        try:
            sp.adjustments[0] = 0.55
            sp.adjustments[1] = 0.34
        except Exception:
            pass
        return sp

    def line(self, x, y, w, h, color, lw=1.5, dash=None):
        from pptx.enum.shapes import MSO_CONNECTOR
        cn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                             Inches(x + w), Inches(y + h))
        cn.line.color.rgb = rgb(color)
        cn.line.width = Pt(lw)
        if dash:
            cn.line.dash_style = dash
        return cn

    # -- icons (pure geometry; no font glyphs) -----------------------------
    def icon(self, kind, cx, cy, s, color, on=None):
        """draw icon centered at (cx,cy) inside square of side s; optional halo circle `on`."""
        x, y = cx - s / 2.0, cy - s / 2.0
        if on:
            self.oval(cx - s * 0.82, cy - s * 0.82, s * 1.64, s * 1.64, fill=on)
        k = kind
        if k == "cross":
            self.shape(MSO_SHAPE.CROSS, x + s * 0.18, y + s * 0.18, s * 0.64, s * 0.64, fill=color)
        elif k == "check":
            a = self.rect(cx - s * 0.36, cy - s * 0.07, s * 0.44, s * 0.16, fill=color, rot=45)
            b = self.rect(cx - s * 0.10, cy - s * 0.30, s * 0.62, s * 0.16, fill=color, rot=-45)
        elif k == "alert":
            self.shape(MSO_SHAPE.ISOSCELES_TRIANGLE, x + s * 0.06, y + s * 0.10, s * 0.88, s * 0.80,
                       fill=color)
            self.rect(cx - s * 0.045, cy - s * 0.16, s * 0.09, s * 0.30, fill=WHITE)
            self.oval(cx - s * 0.05, cy + s * 0.20, s * 0.10, s * 0.10, fill=WHITE)
        elif k == "person":
            self.oval(cx - s * 0.18, y + s * 0.06, s * 0.36, s * 0.36, fill=color)
            self.shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x + s * 0.12, cy + s * 0.06,
                       s * 0.76, s * 0.44, fill=color, rad=0.5)
        elif k == "bed":
            self.rect(x + s * 0.04, cy + s * 0.06, s * 0.92, s * 0.14, fill=color)          # frame
            self.rect(x + s * 0.04, y + s * 0.16, s * 0.12, s * 0.58, fill=color)           # headboard
            self.rect(x + s * 0.86, cy - s * 0.06, s * 0.10, s * 0.36, fill=color)          # footboard
            self.rrect(x + s * 0.20, cy - s * 0.20, s * 0.30, s * 0.20, fill=color, rad=0.4)  # pillow
            self.rrect(x + s * 0.20, cy - s * 0.10, s * 0.62, s * 0.18, fill=color, rad=0.4)  # mattress
        elif k == "hospital":
            self.rect(x + s * 0.16, y + s * 0.30, s * 0.68, s * 0.62, fill=color)
            self.rect(cx - s * 0.06, y + s * 0.06, s * 0.12, s * 0.34, fill=color)
            self.rect(cx - s * 0.17, y + s * 0.17, s * 0.34, s * 0.12, fill=color)
            self.rect(cx - s * 0.08, y + s * 0.66, s * 0.16, s * 0.26, fill=on or WHITE)
        elif k == "clock":
            self.oval(x + s * 0.08, y + s * 0.08, s * 0.84, s * 0.84, fill=None, line=color, lw=2.2)
            self.rect(cx - s * 0.035, cy - s * 0.28, s * 0.07, s * 0.30, fill=color)
            self.rect(cx - s * 0.02, cy - s * 0.035, s * 0.26, s * 0.07, fill=color)
        elif k == "lock":
            self.oval(cx - s * 0.24, y + s * 0.04, s * 0.48, s * 0.48, fill=None, line=color, lw=2.4)
            self.rrect(x + s * 0.12, cy - s * 0.02, s * 0.76, s * 0.52, fill=color, rad=0.18)
        elif k == "db":
            self.shape(MSO_SHAPE.CAN, x + s * 0.12, y + s * 0.08, s * 0.76, s * 0.84, fill=color)
        elif k == "gear":
            self.shape(MSO_SHAPE.GEAR_6, x + s * 0.08, y + s * 0.08, s * 0.84, s * 0.84, fill=color)
        elif k == "pin":
            self.oval(cx - s * 0.26, y + s * 0.06, s * 0.52, s * 0.52, fill=color)
            self.shape(MSO_SHAPE.ISOSCELES_TRIANGLE, cx - s * 0.22, cy - s * 0.02, s * 0.44,
                       s * 0.46, fill=color, rot=180)
        elif k == "doc":
            self.shape(MSO_SHAPE.FOLDED_CORNER, x + s * 0.16, y + s * 0.06, s * 0.68, s * 0.88,
                       fill=color)
        elif k == "chart":
            self.rect(x + s * 0.12, cy + s * 0.02, s * 0.18, s * 0.40, fill=color)
            self.rect(x + s * 0.41, y + s * 0.18, s * 0.18, s * 0.74, fill=color)
            self.rect(x + s * 0.70, cy - s * 0.14, s * 0.18, s * 0.56, fill=color)
        elif k == "users":
            self.oval(x + s * 0.10, y + s * 0.10, s * 0.30, s * 0.30, fill=color)
            self.shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x + s * 0.02, cy + s * 0.02,
                       s * 0.46, s * 0.40, fill=color, rad=0.5)
            self.oval(x + s * 0.56, y + s * 0.18, s * 0.26, s * 0.26, fill=color)
            self.shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x + s * 0.50, cy + s * 0.10,
                       s * 0.40, s * 0.34, fill=color, rad=0.5)
        elif k == "bolt":
            self.shape(MSO_SHAPE.LIGHTNING_BOLT, x + s * 0.14, y + s * 0.06, s * 0.72, s * 0.88,
                       fill=color)
        elif k == "route":
            self.oval(x + s * 0.04, y + s * 0.56, s * 0.24, s * 0.24, fill=color)
            self.oval(cx - s * 0.12, cy - s * 0.12, s * 0.24, s * 0.24, fill=color)
            self.oval(x + s * 0.72, y + s * 0.16, s * 0.24, s * 0.24, fill=color)
            self.line(x + s * 0.24, y + s * 0.68, s * 0.36, -s * 0.36, color, lw=2.0)
            self.line(cx, cy, s * 0.36, -s * 0.24, color, lw=2.0)
        elif k == "flow":
            self.shape(MSO_SHAPE.CHEVRON, x + s * 0.02, y + s * 0.22, s * 0.46, s * 0.56, fill=color)
            self.shape(MSO_SHAPE.CHEVRON, x + s * 0.42, y + s * 0.22, s * 0.46, s * 0.56, fill=color)
        elif k == "list":
            for i in range(3):
                yy = y + s * (0.14 + 0.28 * i)
                self.oval(x + s * 0.08, yy + s * 0.03, s * 0.12, s * 0.12, fill=color)
                self.rect(x + s * 0.30, yy + s * 0.05, s * 0.60, s * 0.09, fill=color)
        elif k == "star":
            self.shape(MSO_SHAPE.STAR_5_POINT, x + s * 0.08, y + s * 0.08, s * 0.84, s * 0.84,
                       fill=color)
        elif k == "shieldcheck":
            self.shape(MSO_SHAPE.PENTAGON, x + s * 0.10, y + s * 0.06, s * 0.80, s * 0.62,
                       fill=color, rot=180)
            self.rect(cx - s * 0.26, cy - s * 0.10, s * 0.30, s * 0.11, fill=on or WHITE, rot=45)
            self.rect(cx - s * 0.08, cy - s * 0.26, s * 0.42, s * 0.11, fill=on or WHITE, rot=-45)
        elif k == "refresh":
            self.shape(MSO_SHAPE.BLOCK_ARC, x + s * 0.10, y + s * 0.10, s * 0.80, s * 0.80,
                       fill=color, rot=30)
            self.shape(MSO_SHAPE.ISOSCELES_TRIANGLE, x + s * 0.52, y + s * 0.02, s * 0.30, s * 0.26,
                       fill=color, rot=90)
        elif k == "eye":
            self.shape(MSO_SHAPE.OVAL, x + s * 0.06, y + s * 0.24, s * 0.88, s * 0.52, fill=color)
            self.oval(cx - s * 0.14, cy - s * 0.14, s * 0.28, s * 0.28, fill=on or WHITE)
        elif k == "grid":
            for i in range(2):
                for j in range(2):
                    self.rrect(x + s * (0.08 + 0.46 * j), y + s * (0.08 + 0.46 * i),
                               s * 0.38, s * 0.38, fill=color, rad=0.2)
        else:
            self.oval(x + s * 0.15, y + s * 0.15, s * 0.7, s * 0.7, fill=color)
        return None

    # -- page furniture ----------------------------------------------------
    def header(self, kicker, kcolor, ktext_color, title, sub=None, dark=False,
               title_size=29, y=0.42):
        self.pill(M, y, 0.10, 0.30, fill=kcolor)
        self.text(M + 0.20, y + 0.015, 8.6, 0.30,
                  [{"t": kicker, "size": 10.5, "color": kcolor if not dark else kcolor,
                    "bold": True}], anchor="middle")
        tcol = WHITE if dark else ELECTRIC
        self.text(M, y + 0.40, 12.2, 0.62,
                  [{"t": title, "size": title_size, "color": tcol, "bold": True, "font": HEAD}],
                  anchor="top", name=title)
        top = y + 1.06
        if sub:
            self.text(M, top, 11.6, 0.34,
                      [{"t": sub, "size": 12.5, "color": (WHITE if dark else INK_SOFT)}],
                      name="sub")
            top += 0.40
        return top

    def footer(self, dark=False):
        c = WHITE if dark else "8A90A6"
        self.text(M, 7.10, 8.0, 0.26,
                  [{"t": "CareRoute — Emergency Bed Coordination System   ·   Team Tech Head",
                    "size": 8.5, "color": c}], check=False)
        self.text(11.35, 7.10, 1.43, 0.26,
                  [{"t": f"{self._n:02d} / 10", "size": 8.5, "color": c, "bold": True,
                    "align": "right"}], check=False)

    def save(self, path):
        self.prs.save(path)
